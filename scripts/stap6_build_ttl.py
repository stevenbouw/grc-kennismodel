#!/usr/bin/env python3
"""Stap 6 v4.5.0 builder: CSF Reference Tool IRs → SKOS in target-modules.

Filter: 3 source-prefixes uit csf2.xlsx K5 ('CSF 2.0' sheet):
  - 'SP 800-53 Rev 5.2.0'                  → m11-nist-800-53.ttl
  - 'ISO/IEC 27001:2022: Annex A Controls' → m08-bio20.ttl (D5 sameAs-brug)
  - 'ISO/IEC 27001:2022: Mandatory Clause' → m09-iso27001-ext.ttl

Match-type uniform: skos:closeMatch (Aanpassing 2).
Bron-attribuering: blok-comment (Aanpassing 3), verwijst naar ext:Attr_NIST_CSF_2_0_Reference_Tool_2026.

CSF-niveau-mappings (Function + Category + Subcategory) allen meegenomen.
"""
import openpyxl
import re
import json
from collections import Counter, defaultdict
from pathlib import Path
from rdflib import Graph, Namespace, RDF, OWL, URIRef

# === Inlees + state-machine ===
wb = openpyxl.load_workbook("/mnt/user-data/uploads/csf2.xlsx", data_only=True)
ws = wb['CSF 2.0']

func_re = re.compile(r'^([A-Z]+) \(([A-Z]{2})\):')
cat_re = re.compile(r'^([^(]+?) \(([A-Z]{2}\.[A-Z]{2})\):')
sub_re = re.compile(r'^([A-Z]{2}\.[A-Z]{2}-\d{2,3}):')

TARGET_PREFIXES = {
    "SP 800-53 Rev 5.2.0":                    "sp800_53",
    "ISO/IEC 27001:2022: Annex A Controls":   "iso_annex",
    "ISO/IEC 27001:2022: Mandatory Clause":   "iso_clause",
}

CSF_FUNC_NAME = {"GV":"GOVERN","ID":"IDENTIFY","PR":"PROTECT","DE":"DETECT","RS":"RESPOND","RC":"RECOVER"}
CAT_IRI_SUFFIX = {  # uit Stap 3
    "GV.OC":"OrganizationalContext","GV.RM":"RiskManagementStrategy","GV.RR":"RolesResponsibilitiesAndAuthorities",
    "GV.PO":"Policy","GV.OV":"Oversight","GV.SC":"CybersecuritySupplyChainRiskManagement",
    "ID.AM":"AssetManagement","ID.RA":"RiskAssessment","ID.IM":"Improvement",
    "PR.AA":"IdentityManagementAuthenticationAndAccessControl","PR.AT":"AwarenessAndTraining",
    "PR.DS":"DataSecurity","PR.PS":"PlatformSecurity","PR.IR":"TechnologyInfrastructureResilience",
    "DE.CM":"ContinuousMonitoring","DE.AE":"AdverseEventAnalysis",
    "RS.MA":"IncidentManagement","RS.AN":"IncidentAnalysis","RS.CO":"IncidentResponseReportingAndCommunication","RS.MI":"IncidentMitigation",
    "RC.RP":"IncidentRecoveryPlanExecution","RC.CO":"IncidentRecoveryCommunication",
}

def csf_subject_iri(code: str, level: str) -> str:
    """level ∈ {'F','C','S'}"""
    if level == "F":
        return f"csf:{CSF_FUNC_NAME[code]}"
    if level == "C":
        suffix = CAT_IRI_SUFFIX[code]
        return f"csf:{code.split('.')[0]}_{suffix}"
    if level == "S":
        return f"csf:{code.replace('.', '_').replace('-', '_')}"
    return None


def resolve_target(prefix_key: str, ref: str):
    """Return (target_iri_or_None, log_status)."""
    if prefix_key == "sp800_53":
        # 'AC-01' → strip leading zero → NIST_AC_1; 'CM-07(02)' → enhancement → None (G1)
        if "(" in ref:
            return None, "sp_enhancement_not_in_m11"
        m = re.match(r'^([A-Z]+)-(\d+)$', ref)
        if m:
            fam, num = m.groups()
            return f"ext:NIST_{fam}_{int(num)}", "ok"
        return None, "sp_format_unknown"
    
    if prefix_key == "iso_annex":
        # '5.1' → bio:ISO27002_5_01 (zfill on minor); '5.10' → bio:ISO27002_5_10
        m = re.match(r'^(\d+)\.(\d+)$', ref)
        if m:
            x, y = m.groups()
            return f"bio:ISO27002_{x}_{y.zfill(2)}", "ok"
        return None, "iso_annex_format_unknown"
    
    if prefix_key == "iso_clause":
        # '6.1' → ext:ISO27001_6_1; '6.1.2' → _6_1_2; '4.2(a)' → drop letter → _4_2; '6.1,' → strip → _6_1
        ref = ref.rstrip(",").strip()
        # Drop sub-paragraph parens: '4.2(a)' or '4.2 (a)' → '4.2'
        ref = re.sub(r'\s*\([a-z]\)\s*$', '', ref)
        m = re.match(r'^(\d+)(?:\.(\d+))?(?:\.(\d+))?$', ref)
        if m:
            parts = [g for g in m.groups() if g is not None]
            return "ext:ISO27001_" + "_".join(parts), "ok"
        return None, "iso_clause_format_unknown"
    
    return None, "prefix_unknown"


# Bouw target-IRI-existence-set
g_check = Graph()
g_check.parse("m11-nist-800-53.ttl", format="turtle")
g_check.parse("m09-iso27001-ext.ttl", format="turtle")
g_check.parse("m08-bio20.ttl", format="turtle")
existing_iris = {str(s) for s in g_check.subjects(RDF.type, OWL.NamedIndividual)}

# Bouw csf-subject-existence-set
g_csf = Graph(); g_csf.parse("m21-csf.ttl", format="turtle")
CSF = Namespace("https://grc.example.org/csf/")
existing_csf = {str(s) for s in g_csf.subjects(RDF.type, OWL.NamedIndividual) if str(s).startswith("https://grc.example.org/csf/")}


# === Parse csf2.xlsx ===
current_F, current_C, current_S = None, None, None
mappings_raw = []  # (csf_subject_iri, prefix_key, ref_normalized, target_iri)
unresolved = []
prefix_log = Counter()  # totaal per prefix
sub_paragraph_log = []  # voor patch-rapport
trailing_comma_log = []

for r in range(2, ws.max_row + 1):
    k1 = ws.cell(row=r, column=1).value
    k2 = ws.cell(row=r, column=2).value
    k3 = ws.cell(row=r, column=3).value
    k5 = ws.cell(row=r, column=5).value
    
    # Update tracking (alleen indien regex daadwerkelijk matcht)
    new_F = new_C = new_S = False
    if k1:
        m = func_re.match(str(k1).strip())
        if m: 
            current_F = m.group(2); current_C = None; current_S = None
            new_F = True
    if k2:
        m = cat_re.match(str(k2).strip())
        if m: 
            current_C = m.group(2); current_S = None
            new_C = True
    if k3:
        m = sub_re.match(str(k3).strip())
        if m: 
            current_S = m.group(1)
            new_S = True
    
    # Bepaal level: meest-specifieke nieuw header heeft voorrang; anders most-specific current state
    if new_S:
        level_now = "S"; level_code = current_S
    elif new_C:
        level_now = "C"; level_code = current_C
    elif new_F:
        level_now = "F"; level_code = current_F
    else:
        if current_S: level_now = "S"; level_code = current_S
        elif current_C: level_now = "C"; level_code = current_C
        elif current_F: level_now = "F"; level_code = current_F
        else: level_now = None; level_code = None
    
    if level_now is None: continue
    if not k5: continue
    csf_iri_str = csf_subject_iri(level_code, level_now)
    if csf_iri_str is None:
        continue
    
    for line in str(k5).split("\n"):
        line = line.strip()
        if not line: continue
        # Find matching target-prefix
        prefix_key = None
        for prefix, key in TARGET_PREFIXES.items():
            if line.startswith(prefix + ":"):
                ref = line[len(prefix)+1:].strip()
                prefix_key = key
                break
        if prefix_key is None: continue
        prefix_log[prefix_key] += 1
        
        # Track edge cases
        if prefix_key == "iso_clause" and "(" in ref:
            sub_paragraph_log.append((csf_iri_str, ref))
        if prefix_key == "iso_clause" and ref.rstrip(",").strip() != ref.strip():
            trailing_comma_log.append((csf_iri_str, ref))
        
        target_iri, status = resolve_target(prefix_key, ref)
        if status != "ok":
            unresolved.append((csf_iri_str, prefix_key, ref, status))
            continue
        
        # Verifieer target-IRI bestaat
        if target_iri.startswith("ext:"):
            iri_full = "https://grc.example.org/extended/" + target_iri[4:]
        elif target_iri.startswith("bio:"):
            iri_full = "https://grc.example.org/bio/" + target_iri[4:]
        else:
            unresolved.append((csf_iri_str, prefix_key, ref, "unknown-namespace"))
            continue
        
        if iri_full not in existing_iris:
            unresolved.append((csf_iri_str, prefix_key, ref, f"target-missing: {target_iri}"))
            continue
        
        # Verifieer CSF-subject bestaat
        csf_iri_full = "https://grc.example.org/csf/" + csf_iri_str[4:]
        if csf_iri_full not in existing_csf:
            unresolved.append((csf_iri_str, prefix_key, ref, f"csf-subject-missing: {csf_iri_str}"))
            continue
        
        mappings_raw.append({
            "csf_subject": csf_iri_str,
            "csf_level": level_now,
            "prefix_key": prefix_key,
            "ref": ref,
            "target": target_iri,
        })

# Dedup per (csf_subject, target)
unique_pairs = set()
unique_mappings = []
for m in mappings_raw:
    key = (m["csf_subject"], m["target"])
    if key in unique_pairs: continue
    unique_pairs.add(key)
    unique_mappings.append(m)

print(f"=== Statistieken ===")
print(f"Raw mappings na filter:    {len(mappings_raw)}")
print(f"Unieke (csf,target)-paren: {len(unique_pairs)}")
print(f"Unresolved:                {len(unresolved)}")
dedup_pct = (len(mappings_raw)-len(unique_pairs))/len(mappings_raw)*100 if mappings_raw else 0
print(f"Dedup-reductie:            {dedup_pct:.1f}%")

print(f"\n=== Verdeling per prefix ===")
for k, c in prefix_log.most_common():
    print(f"  {k}: {c} raw")

per_prefix_unique = Counter(m["prefix_key"] for m in unique_mappings)
print(f"\n=== Unieke per prefix ===")
for k, c in per_prefix_unique.most_common():
    print(f"  {k}: {c} unique")

per_level_unique = Counter(m["csf_level"] for m in unique_mappings)
print(f"\n=== Unieke per CSF-level ===")
for k, c in per_level_unique.most_common():
    print(f"  Level {k}: {c}")

print(f"\n=== Edge cases ===")
print(f"  Sub-paragraph (parens) in iso_clause: {len(sub_paragraph_log)} (mapped naar parent-clause)")
print(f"  Trailing comma in iso_clause: {len(trailing_comma_log)}")
unresolved_by_status = Counter(u[3] if not u[3].startswith('target-missing') and not u[3].startswith('csf-subject') else u[3].split(":")[0] for u in unresolved)
print(f"\n  Unresolved-categorieën:")
for k, c in unresolved_by_status.most_common():
    print(f"    {k}: {c}")

# Save voor verdere stappen
mappings_by_prefix = defaultdict(list)
for m in unique_mappings:
    mappings_by_prefix[m["prefix_key"]].append(m)

# === TTL-blokken bouwen ===
def csf_iri_to_target_module_block(prefix_key, target_module_filename, mappings):
    """Bouw TTL-blok-string voor target-module."""
    by_csf = defaultdict(list)
    for m in mappings:
        by_csf[m["csf_subject"]].append(m["target"])
    
    pretty = {
        "sp800_53": ("SP 800-53 Rev 5.2.0 → CSF 2.0", "csf:<CSFComponent> skos:closeMatch ext:NIST_<XX>_<N>"),
        "iso_annex": ("ISO/IEC 27001:2022 Annex A Controls → CSF 2.0", "csf:<CSFComponent> skos:closeMatch bio:ISO27002_<X>_<YY zfill> (via D5 sameAs-brug naar ctrl:)"),
        "iso_clause": ("ISO/IEC 27001:2022 Mandatory Clause → CSF 2.0", "csf:<CSFComponent> skos:closeMatch ext:ISO27001_<X>[_<Y>[_<Z>]]"),
    }
    name, pattern = pretty[prefix_key]
    
    lines = [f"""
# =============================================================================
# v4.5.0 Fase 3 — Stap 6: SKOS-mappings {name}
# Datum: 2026-05-19
# Bron: NIST CSF 2.0 Reference Tool export csf2.xlsx (gegenereerd 2026-05-19)
#       Bron-attribuering: ext:Attr_NIST_CSF_2_0_Reference_Tool_2026 (Public Domain)
# Match-type: skos:closeMatch (Aanpassing 2 — uniform; geen relationship-types in bron)
# Patroon: {pattern}
# =============================================================================
# Bron-attribuering via blok-comment (consistent met v4.4.0 Stap 5 + v4.5.0 Stap 5).
# Geen per-triple ext:sourceAttribution-link.
#
# Aantal unieke (csf,target)-paren in dit blok: {len(mappings)}
# Verdeling CSF-niveau: F={sum(1 for m in mappings if m['csf_level']=='F')}, """
             + f"""C={sum(1 for m in mappings if m['csf_level']=='C')}, S={sum(1 for m in mappings if m['csf_level']=='S')}
"""]
    
    lines.append(f"\n# --- {len(mappings)} skos:closeMatch (gegroepeerd per CSF-subject) ---\n")
    for csf_iri in sorted(by_csf.keys()):
        targets = sorted(set(by_csf[csf_iri]))
        targets_formatted = " ,\n        ".join(targets)
        lines.append(f"""{csf_iri} skos:closeMatch
        {targets_formatted} .
""")
    return "\n".join(lines)


# Append per target-module
target_files = {
    "sp800_53":   "/home/claude/v433/m11-nist-800-53.ttl",
    "iso_annex":  "/home/claude/v433/m08-bio20.ttl",
    "iso_clause": "/home/claude/v433/m09-iso27001-ext.ttl",
}

for prefix_key, fp in target_files.items():
    mappings = mappings_by_prefix[prefix_key]
    block = csf_iri_to_target_module_block(prefix_key, fp, mappings)
    p = Path(fp)
    current = p.read_text()
    p.write_text(current + block)
    print(f"\n✓ Block appended to {p.name} ({len(block):,} chars, {len(mappings)} mappings)")

# Save log
Path("/home/claude/csf_step6_log.json").write_text(json.dumps({
    "unresolved": unresolved,
    "sub_paragraph_log": sub_paragraph_log,
    "trailing_comma_log": trailing_comma_log,
    "stats": {
        "raw_mappings": len(mappings_raw),
        "unique_pairs": len(unique_pairs),
        "dedup_pct": round(dedup_pct, 1),
        "per_prefix_unique": dict(per_prefix_unique),
        "per_level_unique": dict(per_level_unique),
    }
}, indent=2, ensure_ascii=False))
print("\n✓ Saved log")
